"""Excel export in the "discom blocks 0900 / 1700 / intraday" layout.

One sheet per DISCOM plus ODISHA_SUM (the sum of the four DISCOMs), for the selected period:

  Row 1  title, e.g. "TPNODL · October 2025 · 15-min blocks: actual, day-ahead and intra-day forecasts"
  Row 2  "MAPE" with the period MAPE above each "abs % error" column
  Row 3  headers
  Row 4+ one row per 15-minute block

Error = forecast − actual (MW). Abs % error = |forecast − actual| ÷ actual.
MAPE = average of the abs % errors over blocks that have both an actual and a forecast.
"""

from __future__ import annotations

import io
from calendar import monthrange
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import ALL
from .report import build_report, list_discoms

# Sheet names as in the reference workbook (the southern DISCOM may be stored as TPSOSDL)
SHEET_NAME = {"TPSOSDL": "TPSODL"}
SUM_SHEET = "ODISHA_SUM"
REV_COLS = [("DA1", "Day-ahead 09:00"), ("DA2", "Day-ahead 17:00"), ("ID", "Intra-day")]

HEADERS = ["Date", "Block", "Time", "Actual (MW)"]
for _, name in REV_COLS:
    HEADERS += [f"{name} forecast (MW)", f"{name} error (MW)", f"{name} abs % error"]
WIDTHS = [12, 7, 7, 12] + [12, 12, 12] * 3

MW = "#,##0.0"
PCT = "0.00%"
HEAD_FILL = PatternFill("solid", fgColor="E7EBF0")
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFC7D1")


def period_text(start: date, end: date) -> str:
    if start.day == 1 and end == date(start.year, start.month, monthrange(start.year, start.month)[1]):
        return start.strftime("%B %Y")
    if start == end:
        return start.strftime("%d %b %Y")
    return f"{start.strftime('%d %b %Y')} – {end.strftime('%d %b %Y')}"


def file_name(start: date, end: date) -> str:
    if start.day == 1 and end == date(start.year, start.month, monthrange(start.year, start.month)[1]):
        tag = start.strftime("%Y-%m")
    elif start == end:
        tag = start.isoformat()
    else:
        tag = f"{start.isoformat()}_to_{end.isoformat()}"
    return f"odisha-{tag}-discom-blocks-0900-1700-intraday.xlsx"


def _write_sheet(ws, title: str, rep: dict) -> None:
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value="MAPE").font = BOLD

    mape = {s["revision"]: s["mape"] for s in rep["summary"]}
    for i, (code, _) in enumerate(REV_COLS):
        col = 5 + i * 3 + 2  # abs % error column
        if mape.get(code) is not None:
            c = ws.cell(row=2, column=col, value=mape[code] / 100)
            c.number_format = PCT
            c.font = BOLD

    for j, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="left" if j <= 3 else "center", vertical="bottom")
        c.border = Border(bottom=THIN)
    ws.row_dimensions[3].height = 62

    r = 4
    for row in rep["series"]:
        ts = datetime.strptime(row["ts"], "%Y-%m-%d %H:%M")
        a = row["actual"]
        values = [ts.date(), ts.hour * 4 + ts.minute // 15 + 1, ts.strftime("%H:%M"), a]
        for code, _ in REV_COLS:
            f = row.get(code)
            err = round(f - a, 1) if f is not None and a is not None else None
            ape = abs(f - a) / a if f is not None and a else None
            values += [f, err, ape]
        for j, v in enumerate(values, start=1):
            c = ws.cell(row=r, column=j, value=v)
            if j == 1:
                c.number_format = "yyyy-mm-dd"
                c.alignment = Alignment(horizontal="left")
            elif j == 3:
                c.alignment = Alignment(horizontal="left")
            elif j >= 4:
                c.number_format = PCT if (j - 5) % 3 == 2 and j >= 5 else MW
        r += 1

    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"


def build_workbook(start: date, end: date) -> tuple[bytes, str]:
    discoms = list_discoms()
    period = period_text(start, end)
    wb = Workbook()
    wb.remove(wb.active)

    for d in sorted(discoms, key=lambda x: SHEET_NAME.get(x, x)):
        rep = build_report(d, ALL, start, end)
        name = SHEET_NAME.get(d, d)
        _write_sheet(wb.create_sheet(name), f"{name} · {period} · 15-min blocks: actual, day-ahead and intra-day forecasts", rep)

    rep = build_report(ALL, ALL, start, end)
    _write_sheet(wb.create_sheet(SUM_SHEET),
                 f"Odisha, sum of the four DISCOMs · {period} · 15-min blocks: actual, day-ahead and intra-day forecasts", rep)
    wb.active = len(wb.sheetnames) - 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), file_name(start, end)
